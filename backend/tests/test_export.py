"""
test_export.py — tests PowerPoint (.pptx) and Excel (.xlsx) export endpoints.
Asserts key cells:
  - Grand Total: 120,511,647.51 and 1,167
  - Approval counts: 748 / 877 / 31 / 1,411 / 19
  - Forecast category counts: Closed (1245), Commit (488), Best Case (547), Pipeline (538), Blank (268)
"""
import io
import pytest
from pptx import Presentation
import openpyxl
from fastapi.testclient import TestClient

from backend.main import app
from backend.database import get_db
from backend.models.snapshot import UploadSnapshot
from sqlalchemy.orm import sessionmaker


@pytest.fixture
def client(seeded_session, db_engine):
    Session = sessionmaker(bind=db_engine)

    def _override_get_db():
        db = Session()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = _override_get_db
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def test_pptx_export_oct5_key_cells(client, seeded_session):
    snap = seeded_session.query(UploadSnapshot).filter_by(is_active_today=True).first()
    assert snap is not None

    res = client.get(f"/api/export/pptx?snapshot_id={snap.id}")
    assert res.status_code == 200, res.text
    assert "application/vnd.openxmlformats-officedocument.presentationml.presentation" in res.headers["content-type"]
    assert "attachment; filename=" in res.headers["content-disposition"]

    # Parse PPTX from bytes
    prs = Presentation(io.BytesIO(res.content))
    assert len(prs.slides) == 24

    # 1. Slide 2: Expiry Pivot Table -> Grand Total: 120,511,647.51 and 1,167
    s2 = prs.slides[1]
    t2 = [sh for sh in s2.shapes if sh.name.startswith("Expiry pivot table")][0].table
    last_row_idx = len(t2.rows) - 1
    gt_label = t2.cell(last_row_idx, 0).text_frame.text
    gt_acv = t2.cell(last_row_idx, 2).text_frame.text
    gt_cnt = t2.cell(last_row_idx, 3).text_frame.text

    assert gt_label == "Grand Total"
    assert "120,511,647.51" in gt_acv
    assert "1,167" in gt_cnt or "1167" in gt_cnt

    # 2. Slide 3: Approval Status Table -> 748 / 877 / 31 / 1,411 / 19
    s3 = prs.slides[2]
    t3 = [sh for sh in s3.shapes if sh.name.startswith("Approval status table")][0].table
    row1 = [t3.cell(1, c).text_frame.text.replace(",", "") for c in range(1, 7)]
    # Total, Approved, Approved - 2nd, Pending-Approval, Blank, Rejected
    assert row1[0] == "3086"
    assert row1[1] == "748"
    assert row1[2] == "877"
    assert row1[3] == "31"
    assert row1[4] == "1411"
    assert row1[5] == "19"

    # 3. Slide 4: Forecast Category Table -> counts
    s4 = prs.slides[3]
    t4 = [sh for sh in s4.shapes if sh.name.startswith("Forecast category table")][0].table
    cat_counts = {}
    for r in range(1, 6):
        cat = t4.cell(r, 0).text_frame.text
        cnt = t4.cell(r, 1).text_frame.text.replace(",", "")
        cat_counts[cat] = cnt

    assert cat_counts["Closed"] == "1245"
    assert cat_counts["Commit"] == "488"
    assert cat_counts["Best Case"] == "547"
    assert cat_counts["Pipeline"] == "538"
    assert cat_counts["(Blank)"] == "268"


def test_excel_export_sheets_and_values(client, seeded_session):
    snap = seeded_session.query(UploadSnapshot).filter_by(is_active_today=True).first()
    assert snap is not None

    res = client.get(f"/api/export/excel?snapshot_id={snap.id}")
    assert res.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res.headers["content-type"]
    assert "attachment; filename=" in res.headers["content-disposition"]

    # Parse Excel workbook from bytes
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    expected_sheets = [
        "Expiry Q3 Summary",
        "Approval Status Summary",
        "Approval Status BU Summary",
        "Forecast Category Summary",
        "Forecast Movement",
        "ACV Changes",
        "Top 10 Region Summary",
        "Top 10 Region BU Summary",
        "Change Log",
        "All Opportunities",
    ]
    assert wb.sheetnames == expected_sheets

    # Verify Expiry Q3 Summary sheet
    ws1 = wb["Expiry Q3 Summary"]
    # Check headers
    assert ws1.cell(1, 1).value == "Service Expiry Period"
    assert ws1.cell(1, 3).value == "Today Amount"
    # Find Grand Total row
    gt_found = False
    for r in range(2, ws1.max_row + 1):
        if ws1.cell(r, 1).value == "Grand Total":
            assert round(float(ws1.cell(r, 3).value), 2) == 120511647.51
            assert int(ws1.cell(r, 4).value) == 1167
            gt_found = True
            break
    assert gt_found, "Grand Total row not found in Expiry Q3 Summary"

    # Verify Approval Status Summary sheet
    ws2 = wb["Approval Status Summary"]
    assert ws2.cell(1, 1).value == "Opportunity Approval Status"
    app_data = {}
    for r in range(2, ws2.max_row + 1):
        status = ws2.cell(r, 1).value
        count = ws2.cell(r, 2).value
        if status:
            app_data[status] = count
    assert app_data.get("Approved") == 748
    assert app_data.get("Approved - 2nd") == 877
    assert app_data.get("Pending-Approval") == 31
    assert app_data.get("Blank") == 1411
    assert app_data.get("Rejected") == 19
    assert app_data.get("Grand Total") == 3086


def test_export_direct_routes(client):
    # Both /export/... and /api/export/... should be reachable
    res_pptx = client.get("/export/pptx")
    assert res_pptx.status_code == 200
    res_excel = client.get("/export/excel")
    assert res_excel.status_code == 200


def test_pptx_regression_against_template(client, seeded_session):
    snap = seeded_session.query(UploadSnapshot).filter_by(is_active_today=True).first()
    assert snap is not None

    res = client.get(f"/api/export/pptx?snapshot_id={snap.id}")
    assert res.status_code == 200

    prs = Presentation(io.BytesIO(res.content))
    template_path = "backend/templates/Mobileum_Renewals_Template.pptx"
    prs_template = Presentation(template_path)

    # 1. Slide 2 KPI cards: exactly 2 paragraphs, label and single value line
    s2 = prs.slides[1]
    cards = [sh for sh in s2.shapes if sh.name.startswith("KPI card")]
    assert len(cards) == 3
    for c in cards:
        non_empty_paras = [p.text.strip() for p in c.text_frame.paragraphs if p.text.strip()]
        assert len(non_empty_paras) == 2, f"Expected 2 lines (label and value), got: {non_empty_paras}"

    # 2. Slide 2 pivot: Q2-2027 parent row equals sum of child rows
    t2 = [sh for sh in s2.shapes if sh.name.startswith("Expiry pivot table")][0].table
    q2_2027_parent = None
    q2_2027_children = []
    current_q = None

    for r in range(1, len(t2.rows) - 1):
        q_label = t2.cell(r, 0).text.strip()
        cat_label = t2.cell(r, 1).text.strip()
        if q_label:
            current_q = q_label
            if q_label == "Q2-2027":
                q2_2027_parent = {
                    "acv": float(t2.cell(r, 2).text.replace(",", "")),
                    "cnt": int(t2.cell(r, 3).text),
                    "tlw_acv": float(t2.cell(r, 6).text.replace(",", "")),
                    "tlw_cnt": int(t2.cell(r, 7).text),
                }
        elif current_q == "Q2-2027" and cat_label:
            q2_2027_children.append({
                "cat": cat_label,
                "acv": float(t2.cell(r, 2).text.replace(",", "") or 0),
                "cnt": int(t2.cell(r, 3).text or 0),
                "tlw_acv": float(t2.cell(r, 6).text.replace(",", "") or 0),
                "tlw_cnt": int(t2.cell(r, 7).text or 0),
            })

    assert q2_2027_parent is not None
    assert len(q2_2027_children) == 2  # Commit and Best Case
    child_sum_acv = sum(c["acv"] for c in q2_2027_children)
    child_sum_cnt = sum(c["cnt"] for c in q2_2027_children)
    child_sum_tlw_acv = sum(c["tlw_acv"] for c in q2_2027_children)
    child_sum_tlw_cnt = sum(c["tlw_cnt"] for c in q2_2027_children)

    assert round(child_sum_acv, 2) == round(q2_2027_parent["acv"], 2)
    assert child_sum_cnt == q2_2027_parent["cnt"]
    assert round(child_sum_tlw_acv, 2) == round(q2_2027_parent["tlw_acv"], 2)
    assert child_sum_tlw_cnt == q2_2027_parent["tlw_cnt"]

    # 3. Slide 6 #10 opportunity matches template (006Qp00000eV7mAIAS)
    s6 = prs.slides[5]
    t6 = [sh for sh in s6.shapes if sh.name.startswith("Top 10 table")][0].table
    assert t6.cell(10, 1).text.strip() == "006Qp00000eV7mAIAS"

    # 4. Slide 15 BU table is sorted alphabetically
    s15 = prs.slides[14]
    t15 = [sh for sh in s15.shapes if sh.name.startswith("BU approval table")][0].table
    bus = [t15.cell(r, 0).text.strip() for r in range(1, len(t15.rows) - 1)]
    assert bus == sorted(bus, key=lambda s: s.lower())

    # 5. Top 10 tables on slides 6-14 and 16-24 match template opportunities
    for s_idx in range(5, 14):
        s_gen = prs.slides[s_idx]
        s_tpl = prs_template.slides[s_idx]
        t_gen = [sh for sh in s_gen.shapes if sh.name.startswith("Top 10 table")][0].table
        t_tpl = [sh for sh in s_tpl.shapes if sh.name.startswith("Top 10 table")][0].table
        gen_opps = [t_gen.cell(r, 1).text.strip() for r in range(1, 11) if t_gen.cell(r, 1).text.strip()]
        tpl_opps = [t_tpl.cell(r, 1).text.strip() for r in range(1, 11) if t_tpl.cell(r, 1).text.strip()]
        assert gen_opps == tpl_opps, f"Slide {s_idx+1} opp mismatch: {gen_opps} vs {tpl_opps}"

    for s_idx in range(15, 24):
        s_gen = prs.slides[s_idx]
        s_tpl = prs_template.slides[s_idx]
        t_gen = [sh for sh in s_gen.shapes if sh.name.startswith("Top 10 by BU table")][0].table
        t_tpl = [sh for sh in s_tpl.shapes if sh.name.startswith("Top 10 by BU table")][0].table
        gen_opps = [t_gen.cell(r, 2).text.strip() for r in range(1, 11) if t_gen.cell(r, 2).text.strip()]
        tpl_opps = [t_tpl.cell(r, 2).text.strip() for r in range(1, 11) if t_tpl.cell(r, 2).text.strip()]
        assert gen_opps == tpl_opps, f"Slide {s_idx+1} BU opp mismatch: {gen_opps} vs {tpl_opps}"
