"""
PptxExportService — fills Mobileum_Renewals_Template.pptx using database and analytics service.
"""
from __future__ import annotations

import io
import os
from datetime import date
from decimal import Decimal
from pathlib import Path
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from sqlalchemy.orm import Session

from backend.config import settings
from backend.models.snapshot import UploadSnapshot
from backend.services.analytics_service import AnalyticsService, _opps_to_df
from backend.services.context import ADMIN_CONTEXT

GREEN = RGBColor(0x92, 0xD0, 0x6E)
RED = RGBColor(0xF4, 0xB0, 0x84)
BAND = RGBColor(0xDD, 0xEB, 0xF7)


def setc(cell, text: str, fill="keep"):
    """
    Set text of cell's first paragraph run without wiping font formatting.
    """
    p = cell.text_frame.paragraphs[0]
    if p.runs:
        p.runs[0].text = str(text)
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run().text = str(text)
    if fill != "keep":
        if fill is None:
            cell.fill.background()
        else:
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill


def sfill(v, base=None):
    """Return green if positive, red if negative, else base."""
    try:
        x = float(str(v).replace(",", "").replace("$", "").replace("+", "").strip())
    except Exception:
        return base
    return GREEN if x > 0 else RED if x < 0 else base


def money(v, d=2) -> str:
    if v is None:
        return ""
    try:
        x = float(v)
        return f"{x:,.{d}f}"
    except Exception:
        return str(v)


def to_int(v) -> int:
    if v is None:
        return 0
    try:
        return int(float(str(v).replace(",", "")))
    except Exception:
        return 0


def to_float(v) -> float:
    if v is None:
        return 0.0
    try:
        return float(str(v).replace(",", ""))
    except Exception:
        return 0.0


def find(shape_container, prefix: str):
    return [sh for sh in shape_container.shapes if sh.name.startswith(prefix)]


def title(slide) -> str:
    return slide.shapes.title.text_frame.text if slide.shapes.title is not None else ""


def set_text(shape, text: str):
    if shape.text_frame.paragraphs and shape.text_frame.paragraphs[0].runs:
        shape.text_frame.paragraphs[0].runs[0].text = text
        for x in shape.text_frame.paragraphs[0].runs[1:]:
            x.text = ""
    else:
        shape.text_frame.text = text


def set_kpi_card(shape, label: str, value: str):
    """
    Replace existing text in KPI card:
    Paragraph 0: label
    Paragraph 1: one value line
    Clear any extra lines or paragraphs.
    """
    tf = shape.text_frame
    if len(tf.paragraphs) > 0:
        p0 = tf.paragraphs[0]
        if p0.runs:
            p0.runs[0].text = str(label)
            for r in p0.runs[1:]:
                r.text = ""
        else:
            p0.text = str(label)
    if len(tf.paragraphs) > 1:
        p1 = tf.paragraphs[1]
        if p1.runs:
            p1.runs[0].text = str(value)
            for r in p1.runs[1:]:
                r.text = ""
        else:
            p1.text = str(value)
    for p in tf.paragraphs[2:]:
        p.text = ""


def adjust_table_rows(tbl, target_row_count: int):
    """
    Dynamically add or remove rows from a PowerPoint table before the Grand Total row.
    Clones styling of row 2 for added rows.
    """
    import copy
    current_count = len(tbl.rows)
    if current_count < target_row_count:
        template_tr = copy.deepcopy(tbl.rows[2]._tr)
        for tc in template_tr.iter():
            if tc.tag.endswith("}t"):
                tc.text = ""
        while len(tbl.rows) < target_row_count:
            last_idx = len(tbl.rows) - 1
            new_tr = copy.deepcopy(template_tr)
            tbl.rows[last_idx]._tr.addprevious(new_tr)
    elif current_count > target_row_count:
        while len(tbl.rows) > target_row_count:
            remove_idx = len(tbl.rows) - 2
            tbl._tbl.remove(tbl.rows[remove_idx]._tr)


class PptxExportService:
    def __init__(self, db: Session):
        self.db = db
        self.analytics_svc = AnalyticsService(db)

    def generate_pptx(
        self,
        snapshot_id: str | None = None,
        compare_to: str | None = None,
        scope: str = "renewals",
        include_deleted_lost: bool = False,
    ) -> tuple[io.BytesIO, str]:
        # 1. Resolve active snapshot
        snap = self.analytics_svc._resolve_snapshot(ADMIN_CONTEXT, snapshot_id)
        if not snap:
            raise ValueError("No active or specified snapshot found in database.")

        # 2. Resolve comparison snapshot (yesterday by default)
        comp_snap = None
        if compare_to:
            comp_snap = self.analytics_svc.get_snapshot_by_id(ADMIN_CONTEXT, compare_to)
            if not comp_snap:
                # Try date lookup
                try:
                    d = date.fromisoformat(compare_to)
                    comp_snap = self.analytics_svc.get_snapshot_by_date(ADMIN_CONTEXT, d)
                except Exception:
                    pass
        if not comp_snap:
            comp_snap = self.analytics_svc._get_adjacent_snapshot(snap, days_back=1)

        # 3. Resolve last week snapshot (7 days back)
        lw_snap = self.analytics_svc._get_adjacent_snapshot(snap, days_back=7)

        template_path = Path(__file__).resolve().parent.parent / "templates" / "Mobileum_Renewals_Template.pptx"
        if not template_path.exists():
            raise FileNotFoundError(f"Template not found at: {template_path}")

        prs = Presentation(str(template_path))
        fmt_date = snap.snapshot_date.strftime("%d-%b-%Y")

        # Gather data
        opps = self.analytics_svc._opps_for(snap, scope=scope, include_deleted_lost=include_deleted_lost)
        fs = self.analytics_svc.get_forecast_summary(
            ADMIN_CONTEXT, snapshot_id=snap.id, compare_to=comp_snap.id if comp_snap else None,
            scope=scope, include_deleted_lost=include_deleted_lost,
        )
        by_cell = fs.get("by_cell", {})
        by_quarter = fs.get("by_quarter", {})
        totals = fs.get("totals", {})

        comp_data = {}
        if comp_snap:
            comp_data = self.analytics_svc.get_comparison(
                ADMIN_CONTEXT, comp_snap.snapshot_date, snap.snapshot_date,
                scope=scope, include_deleted_lost=include_deleted_lost,
            )

        aps = self.analytics_svc.get_approval_status(
            ADMIN_CONTEXT, snapshot_id=snap.id, compare_to=comp_snap.id if comp_snap else None,
            scope=scope, include_deleted_lost=include_deleted_lost,
        )
        bu_aps = self.analytics_svc.get_approval_by_bu(
            ADMIN_CONTEXT, snapshot_id=snap.id, mode="as_in_excel",
            scope=scope, include_deleted_lost=include_deleted_lost,
        )

        # Load Excel tie-break row order if reference workbook exists
        ra_order: dict[str, int] = {}
        rb_order: dict[str, int] = {}
        excel_summary_path = settings.DATA_DIR / "Renewals Summary 3.xlsx"
        if not excel_summary_path.exists() and snap.source_file_summary:
            excel_summary_path = settings.DATA_DIR / snap.source_file_summary

        if excel_summary_path.exists():
            try:
                rA = pd.read_excel(excel_summary_path, sheet_name="Top 10 Region Summary")
                ra_order = {
                    str(row["Opportunity ID 18 Digit"]).strip(): idx
                    for idx, row in rA.iterrows()
                    if pd.notna(row.get("Opportunity ID 18 Digit"))
                }
            except Exception:
                pass
            try:
                rB = pd.read_excel(excel_summary_path, sheet_name="Top 10 Region BU Summary")
                rb_order = {
                    str(row["Opportunity ID 18 Digit"]).strip(): idx
                    for idx, row in rB.iterrows()
                    if pd.notna(row.get("Opportunity ID 18 Digit"))
                }
            except Exception:
                pass

        for s in prs.slides:
            t = title(s)
            sub = find(s, "Data as of")

            # ---------------------------------------------------------
            # Slide 1: Cover
            # ---------------------------------------------------------
            cov = find(s, "Cover date")
            if cov:
                set_text(cov[0], f"Data as of {fmt_date}")

            # ---------------------------------------------------------
            # Slide 2: Renewals Summary – Expiry Q3
            # ---------------------------------------------------------
            if t.startswith("Renewals Summary"):
                if sub:
                    set_text(
                        sub[0],
                        f"Data as of {fmt_date}   |   T-Y = Today vs Yesterday,  T-LW = Today vs Last Week",
                    )
                tbl_list = find(s, "Expiry pivot table")
                if tbl_list:
                    tbl = tbl_list[0].table

                    # Determine active quarters and categories dynamically
                    canonical_quarters = [
                        "Q1-2026", "Q1-2027", "Q2-2026", "Q2-2027",
                        "Q3-2026", "Q3-2027", "Q4-2026", "Q4-2027",
                    ]
                    canonical_cats = ["Closed", "Commit", "Best Case", "Pipeline"]

                    q_cats = []
                    for q in canonical_quarters:
                        active_cats = []
                        for c in canonical_cats:
                            c_d = by_cell.get(f"{q}|{c}", by_cell.get((q, c), {}))
                            t_c = c_d.get("today", {})
                            y_c = c_d.get("yesterday", {})
                            lw_c = c_d.get("lastweek", {})
                            if (
                                t_c.get("count", 0) != 0 or t_c.get("acv", 0) != 0
                                or y_c.get("count", 0) != 0 or y_c.get("acv", 0) != 0
                                or lw_c.get("count", 0) != 0 or lw_c.get("acv", 0) != 0
                            ):
                                active_cats.append(c)
                        # Fallback for quarters with no deals
                        if not active_cats:
                            active_cats = ["Commit"]
                        q_cats.append((q, active_cats))

                    # 1 header row + sum(1 parent + len(cats)) + 1 Grand Total row
                    total_needed_rows = 1 + sum(1 + len(cats) for _, cats in q_cats) + 1
                    adjust_table_rows(tbl, total_needed_rows)

                    row_idx = 1
                    for q, cats in q_cats:
                        q_d = by_quarter.get(q, {})
                        t_q = q_d.get("today", {})
                        ty_q = q_d.get("delta_vs_yesterday", {})
                        tlw_q = q_d.get("delta_vs_lastweek", {})
                        q_row_vals = [
                            q,
                            " ",
                            money(t_q.get("acv", 0)),
                            str(int(t_q.get("count", 0))),
                            money(ty_q.get("acv", 0)),
                            str(int(ty_q.get("count", 0))),
                            money(tlw_q.get("acv", 0)),
                            str(int(tlw_q.get("count", 0))),
                        ]
                        for j, x in enumerate(q_row_vals):
                            fill = BAND if j < 4 else sfill(x, BAND)
                            setc(tbl.cell(row_idx, j), str(x), fill)
                        row_idx += 1

                        for c in cats:
                            c_d = by_cell.get(f"{q}|{c}", by_cell.get((q, c), {}))
                            t_c = c_d.get("today", {})
                            ty_c = c_d.get("delta_vs_yesterday", {})
                            tlw_c = c_d.get("delta_vs_lastweek", {})
                            c_row_vals = [
                                " ",
                                c,
                                money(t_c.get("acv", 0)),
                                str(int(t_c.get("count", 0))),
                                money(ty_c.get("acv", 0)),
                                str(int(ty_c.get("count", 0))),
                                money(tlw_c.get("acv", 0)),
                                str(int(tlw_c.get("count", 0))),
                            ]
                            for j, x in enumerate(c_row_vals):
                                fill = None if j < 4 else sfill(x, None)
                                setc(tbl.cell(row_idx, j), str(x), fill)
                            row_idx += 1

                    # Grand Total row
                    t_tot = totals.get("today", {})
                    ty_tot = totals.get("delta_vs_yesterday", {})
                    tlw_tot = totals.get("delta_vs_lastweek", {})
                    gt_vals = [
                        "Grand Total",
                        " ",
                        money(to_float(t_tot.get("acv"))),
                        str(to_int(t_tot.get("count"))),
                        money(to_float(ty_tot.get("acv"))),
                        str(to_int(ty_tot.get("count"))),
                        money(to_float(tlw_tot.get("acv"))),
                        str(to_int(tlw_tot.get("count"))),
                    ]
                    for j, x in enumerate(gt_vals):
                        setc(tbl.cell(row_idx, j), str(x), "keep")

                    # 3 KPI cards: clean replacement with label and one value line
                    cards = find(s, "KPI card")
                    t_acv_m = to_float(t_tot.get("acv")) / 1e6
                    ty_acv_m = to_float(ty_tot.get("acv")) / 1e6
                    tlw_acv_m = to_float(tlw_tot.get("acv")) / 1e6
                    ty_sign = "+" if ty_acv_m >= 0 else "-"
                    tlw_sign = "+" if tlw_acv_m >= 0 else "-"

                    if len(cards) >= 3:
                        set_kpi_card(
                            cards[0],
                            "Total ACV (Today)",
                            f"${t_acv_m:,.2f}M  |  {to_int(t_tot.get('count')):,}",
                        )
                        set_kpi_card(
                            cards[1],
                            "vs Yesterday",
                            f"{ty_sign}${abs(ty_acv_m):,.2f}M  |  {to_int(ty_tot.get('count')):+d}",
                        )
                        set_kpi_card(
                            cards[2],
                            "vs Last Week",
                            f"{tlw_sign}${abs(tlw_acv_m):,.2f}M  |  {to_int(tlw_tot.get('count')):+d}",
                        )

            # ---------------------------------------------------------
            # Slide 3: Opportunity Approval Status
            # ---------------------------------------------------------
            elif t.startswith("Opportunity Approval"):
                if sub:
                    set_text(sub[0], f"Data as of {fmt_date}")
                tbl_list = find(s, "Approval status table")
                if tbl_list:
                    tbl = tbl_list[0].table
                    app_map = {r["approval_status"]: r for r in aps.get("rows", [])}
                    order = ["Approved", "Approved - 2nd", "Pending-Approval", "Blank", "Rejected"]
                    counts = [app_map.get(k, {}).get("count", 0) for k in order]
                    acvs = [app_map.get(k, {}).get("acv", 0.0) for k in order]
                    tot_cnt = aps.get("total_count", sum(counts))
                    tot_acv = aps.get("total_acv", sum(acvs))

                    # Row 1: counts
                    cnt_row = ["Opportunities", f"{tot_cnt:,}"] + [f"{c:,}" for c in counts]
                    for j, val in enumerate(cnt_row):
                        setc(tbl.cell(1, j), val)

                    # Row 2: acv
                    acv_row = ["ACV (USD)", f"${tot_acv/1e6:,.2f}M"] + [f"${a/1e6:,.2f}M" for a in acvs]
                    for j, val in enumerate(acv_row):
                        setc(tbl.cell(2, j), val)

                # Chart
                charts = find(s, "Approval chart")
                if charts:
                    ch = charts[0].chart
                    cd = CategoryChartData()
                    cd.categories = ["Approved", "Approved - 2nd", "Pending-Approval", "Blank", "Rejected"]
                    cd.add_series("Opportunities", counts)
                    ch.replace_data(cd)

                # Today's changes
                chg = aps.get("changes", {})
                kp = find(s, "Key points")
                if kp:
                    ps = kp[0].text_frame.paragraphs
                    new_app = chg.get("newly_approved_count", 0)
                    new_pend = chg.get("newly_pending_count", 0)
                    new_rej = chg.get("newly_rejected_count", 0)
                    if len(ps) >= 4:
                        if ps[1].runs:
                            ps[1].runs[0].text = f"Newly approved:  {new_app}"
                            for r in ps[1].runs[1:]:
                                r.text = ""
                        if ps[2].runs:
                            ps[2].runs[0].text = f"Newly pending approval:  {new_pend}"
                            for r in ps[2].runs[1:]:
                                r.text = ""
                        if ps[3].runs:
                            ps[3].runs[0].text = f"Newly rejected:  {new_rej}"
                            for r in ps[3].runs[1:]:
                                r.text = ""

            # ---------------------------------------------------------
            # Slide 4: Forecast Category: Today vs Yesterday
            # ---------------------------------------------------------
            elif t.startswith("Forecast Category"):
                if sub:
                    set_text(sub[0], f"Data as of {fmt_date}")
                tb_list = find(s, "Forecast category table")
                if tb_list:
                    tb = tb_list[0].table
                    fc_list = comp_data.get("forecast_category", [])
                    fc_map = {item["category"]: item for item in fc_list}
                    categories = ["Closed", "Commit", "Best Case", "Pipeline", "Blank"]

                    # Precompute today's values directly from opps for robustness
                    today_fc_counts = {c: 0 for c in categories}
                    today_fc_acv = {c: 0.0 for c in categories}
                    for op in opps:
                        c_norm = op.forecast_category or "Blank"
                        if c_norm in today_fc_counts:
                            today_fc_counts[c_norm] += 1
                            today_fc_acv[c_norm] += float(op.forecast_acv_amount or 0)

                    tot_today_c = sum(fc_map.get(c, {}).get("to_count") if fc_map.get(c, {}).get("to_count") is not None else today_fc_counts[c] for c in categories)
                    tot_today_a = sum(fc_map.get(c, {}).get("to_acv") if fc_map.get(c, {}).get("to_acv") is not None else today_fc_acv[c] for c in categories)
                    tot_yest_a = sum(to_float(fc_map.get(c, {}).get("from_acv")) for c in categories)
                    tot_cnt_diff = sum(to_int(fc_map.get(c, {}).get("count_diff")) for c in categories)
                    tot_acv_diff = sum(to_float(fc_map.get(c, {}).get("acv_diff")) for c in categories)

                    for i, c in enumerate(categories, 1):
                        item = fc_map.get(c, {})
                        c_label = "(Blank)" if c == "Blank" else c
                        c_cnt = item.get("to_count") if item.get("to_count") is not None else today_fc_counts[c]
                        c_acv = item.get("to_acv") if item.get("to_acv") is not None else today_fc_acv[c]
                        row_vals = [
                            c_label,
                            str(to_int(c_cnt)),
                            money(to_float(c_acv), 0),
                            money(to_float(item.get("from_acv")), 0),
                            f"{to_int(item.get('count_diff')):+d}",
                            f"{to_float(item.get('acv_diff')):+,.0f}",
                        ]
                        for j, x in enumerate(row_vals):
                            fill = sfill(x, None) if j >= 4 else "keep"
                            setc(tb.cell(i, j), str(x), fill)

                    # Total row
                    tot_vals = [
                        "Total",
                        str(int(tot_today_c)),
                        money(tot_today_a, 0),
                        money(tot_yest_a, 0),
                        f"{int(tot_cnt_diff):+d}",
                        f"{tot_acv_diff:+,.0f}",
                    ]
                    for j, x in enumerate(tot_vals):
                        fill = sfill(x, BAND) if j >= 4 else BAND
                        setc(tb.cell(6, j), str(x), fill)

                # Movement table
                mt_list = find(s, "Forecast movement table")
                if mt_list:
                    mt = mt_list[0].table
                    movements = comp_data.get("movement", [])
                    for r_idx in range(1, len(mt.rows)):
                        if r_idx - 1 < len(movements):
                            mv = movements[r_idx - 1]
                            from_c = mv.get("from") or "(Blank)"
                            to_c = mv.get("to") or "(Blank)"
                            mv_text = f"{from_c} → {to_c}"
                            setc(mt.cell(r_idx, 0), mv_text)
                            setc(mt.cell(r_idx, 1), str(int(mv.get("count", 0))))
                        else:
                            setc(mt.cell(r_idx, 0), " ")
                            setc(mt.cell(r_idx, 1), " ")

            # ---------------------------------------------------------
            # Slide 5: Top 10 Forecast ACV Changes vs Yesterday
            # ---------------------------------------------------------
            elif t.startswith("Top 10 Forecast ACV"):
                if sub:
                    set_text(sub[0], f"Data as of {fmt_date}")
                tb_list = find(s, "ACV change table")
                if tb_list:
                    tb = tb_list[0].table
                    movers = comp_data.get("biggest_movers", [])[:10]
                    for i in range(10):
                        if i < len(movers):
                            m = movers[i]
                            vals = [
                                i + 1,
                                m.get("opportunity_name") or "",
                                money(m.get("old_acv", 0), 0),
                                money(m.get("new_acv", 0), 0),
                                m.get("tag", "Increase"),
                            ]
                        else:
                            vals = [i + 1, " ", " ", " ", " "]
                        for j, v in enumerate(vals):
                            setc(tb.cell(i + 1, j), str(v))

            # ---------------------------------------------------------
            # Slide 6..14: Top 10 Opportunities – <Region>
            # ---------------------------------------------------------
            elif t.startswith("Top 10 Opportunities –"):
                reg = t.split("–", 1)[1].strip()
                reg_opps = [o for o in opps if o.sub_region == reg]
                reg_acv = sum(float(o.forecast_acv_amount or 0) for o in reg_opps)
                reg_cnt = len(reg_opps)
                if sub:
                    set_text(
                        sub[0],
                        f"Data as of {fmt_date}   |   {reg} total: ${reg_acv/1e6:,.2f}M across {reg_cnt:,} opportunities",
                    )
                tb_list = find(s, "Top 10 table")
                if tb_list:
                    tb = tb_list[0].table
                    top10 = sorted(
                        reg_opps,
                        key=lambda o: (
                            -float(o.forecast_acv_amount or 0),
                            ra_order.get(str(o.opportunity_id_18 or "").strip(), 999999),
                        ),
                    )[:10]
                    for i in range(10):
                        if i < len(top10):
                            o = top10[i]
                            vals = [
                                i + 1,
                                o.opportunity_id_18 or "",
                                o.opportunity_name or "",
                                money(o.forecast_acv_amount),
                            ]
                        else:
                            vals = [i + 1, " ", " ", " "]
                        for j, v in enumerate(vals):
                            setc(tb.cell(i + 1, j), str(v))
                    top10_sum = sum(float(o.forecast_acv_amount or 0) for o in top10)
                    setc(tb.cell(11, 3), money(top10_sum))

            # ---------------------------------------------------------
            # Slide 15: Business Unit – Opportunities & Approval Status
            # ---------------------------------------------------------
            elif t.startswith("Business Unit –"):
                if sub:
                    set_text(
                        sub[0],
                        f"Data as of {fmt_date}   |   Count of opportunities by Business Unit; Blank = no approval status yet",
                    )
                tb_list = find(s, "BU approval table")
                if tb_list:
                    tb = tb_list[0].table
                    bu_rows = sorted(
                        bu_aps.get("as_in_excel_rows", []),
                        key=lambda r: str(r.get("business_unit", "")).lower(),
                    )
                    bu_total = bu_aps.get("total", {})
                    for i in range(1, len(tb.rows) - 1):
                        idx = i - 1
                        if idx < len(bu_rows):
                            r = bu_rows[idx]
                            vals = [
                                r.get("business_unit"),
                                f"{r.get('total_count', 0):,}",
                                str(r.get("approved_count", 0)) if r.get("approved_count", 0) > 0 else "–",
                                str(r.get("approved_2nd_count", 0))
                                if r.get("approved_2nd_count", 0) > 0
                                else "–",
                                str(r.get("pending_count", 0)) if r.get("pending_count", 0) > 0 else "–",
                                str(r.get("blank_count", 0)) if r.get("blank_count", 0) > 0 else "–",
                                str(r.get("rejected_count", 0)) if r.get("rejected_count", 0) > 0 else "–",
                            ]
                        else:
                            vals = [" ", " ", " ", " ", " ", " ", " "]
                        for j, v in enumerate(vals):
                            setc(tb.cell(i, j), str(v))

                    # Total row
                    last_r = len(tb.rows) - 1
                    tot_vals = [
                        "Total",
                        f"{bu_total.get('total_count', 0):,}",
                        f"{bu_total.get('approved_count', 0):,}",
                        f"{bu_total.get('approved_2nd_count', 0):,}",
                        f"{bu_total.get('pending_count', 0):,}",
                        f"{bu_total.get('blank_count', 0):,}",
                        f"{bu_total.get('rejected_count', 0):,}",
                    ]
                    for j, v in enumerate(tot_vals):
                        setc(tb.cell(last_r, j), str(v))

            # ---------------------------------------------------------
            # Slide 16..24: Top 10 by Business Unit – <Region>
            # ---------------------------------------------------------
            elif t.startswith("Top 10 by Business Unit –"):
                reg = t.split("–", 1)[1].strip()
                reg_opps = [o for o in opps if o.sub_region == reg]

                # leading BU determined by single highest ACV opportunity in that region
                bu_max = {}
                for o in reg_opps:
                    bu = o.business_unit_raw or "Unknown"
                    val = float(o.forecast_acv_amount or 0)
                    if bu not in bu_max or val > bu_max[bu]:
                        bu_max[bu] = val

                leading_bu = max(bu_max.keys(), key=lambda b: bu_max[b]) if bu_max else "None"
                bu_opps = [o for o in reg_opps if (o.business_unit_raw or "Unknown") == leading_bu]
                tot_bu_acv = sum(float(o.forecast_acv_amount or 0) for o in bu_opps)
                tot_bu_cnt = len(bu_opps)

                opp_word = "opportunity" if tot_bu_cnt == 1 else "opportunities"
                if sub:
                    set_text(
                        sub[0],
                        f"Data as of {fmt_date}   |   {reg} – {leading_bu}: ${tot_bu_acv/1e6:,.2f}M across {tot_bu_cnt:,} {opp_word}",
                    )

                tb_list = find(s, "Top 10 by BU table")
                if tb_list:
                    tb = tb_list[0].table
                    top10_bu = sorted(
                        bu_opps,
                        key=lambda o: (
                            -float(o.forecast_acv_amount or 0),
                            rb_order.get(str(o.opportunity_id_18 or "").strip(), 999999),
                        ),
                    )[:10]
                    for i in range(10):
                        if i < len(top10_bu):
                            o = top10_bu[i]
                            vals = [
                                i + 1,
                                leading_bu,
                                o.opportunity_id_18 or "",
                                o.opportunity_name or "",
                                money(o.forecast_acv_amount),
                            ]
                        else:
                            vals = [i + 1, " ", " ", " ", " "]
                        for j, v in enumerate(vals):
                            setc(tb.cell(i + 1, j), str(v))
                    top10_bu_sum = sum(float(o.forecast_acv_amount or 0) for o in top10_bu)
                    setc(tb.cell(11, 4), money(top10_bu_sum))

        output = io.BytesIO()
        prs.save(output)
        output.seek(0)
        filename = f"Mobileum_Renewals_Daily_Update_{snap.snapshot_date.isoformat()}.pptx"
        return output, filename
