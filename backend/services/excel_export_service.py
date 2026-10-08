"""
ExcelExportService — builds a formatted multi-sheet Excel workbook using openpyxl.
"""
from __future__ import annotations

import io
from datetime import date
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from backend.models.change_log import ChangeLog
from backend.services.analytics_service import AnalyticsService, _opps_to_df
from backend.services.context import ADMIN_CONTEXT

NAVY_FILL = PatternFill(start_color="12284C", end_color="12284C", fill_type="solid")
BAND_FILL = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
GREEN_FONT = Font(name="Calibri", size=11, color="006100", bold=True)
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
RED_FONT = Font(name="Calibri", size=11, color="9C0006", bold=True)

WHITE_BOLD_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BOLD_FONT = Font(name="Calibri", size=11, bold=True)
REGULAR_FONT = Font(name="Calibri", size=11)

THIN_BORDER = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)


def style_header_row(ws, cols_count: int):
    for col in range(1, cols_count + 1):
        cell = ws.cell(1, col)
        cell.fill = NAVY_FILL
        cell.font = WHITE_BOLD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(cols_count)}1"


def autofit_columns(ws, min_width: int = 12, max_width: int = 45):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(min(max_len + 3, max_width), min_width)


def to_int(v) -> int:
    if v is None:
        return 0
    try:
        return int(float(str(v).replace(",", "")))
    except Exception:
        return 0


def to_float(v) -> float:
    if v is None:
        return 0.0
    try:
        return float(str(v).replace(",", ""))
    except Exception:
        return 0.0


class ExcelExportService:
    def __init__(self, db: Session):
        self.db = db
        self.analytics_svc = AnalyticsService(db)

    def generate_excel(
        self,
        snapshot_id: str | None = None,
        compare_to: str | None = None,
        scope: str = "renewals",
        include_deleted_lost: bool = False,
    ) -> tuple[io.BytesIO, str]:
        # 1. Resolve active snapshot
        snap = self.analytics_svc._resolve_snapshot(ADMIN_CONTEXT, snapshot_id)
        if not snap:
            raise ValueError("No active or specified snapshot found in database.")

        # 2. Resolve comparison snapshot (yesterday by default)
        comp_snap = None
        if compare_to:
            comp_snap = self.analytics_svc.get_snapshot_by_id(ADMIN_CONTEXT, compare_to)
            if not comp_snap:
                try:
                    d = date.fromisoformat(compare_to)
                    comp_snap = self.analytics_svc.get_snapshot_by_date(ADMIN_CONTEXT, d)
                except Exception:
                    pass
        if not comp_snap:
            comp_snap = self.analytics_svc._get_adjacent_snapshot(snap, days_back=1)

        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # remove default sheet

        opps = self.analytics_svc._opps_for(snap, scope=scope, include_deleted_lost=include_deleted_lost)
        fs = self.analytics_svc.get_forecast_summary(
            ADMIN_CONTEXT, snapshot_id=snap.id, compare_to=comp_snap.id if comp_snap else None,
            scope=scope, include_deleted_lost=include_deleted_lost,
        )
        aps = self.analytics_svc.get_approval_status(
            ADMIN_CONTEXT, snapshot_id=snap.id, compare_to=comp_snap.id if comp_snap else None,
            scope=scope, include_deleted_lost=include_deleted_lost,
        )
        bu_aps = self.analytics_svc.get_approval_by_bu(
            ADMIN_CONTEXT, snapshot_id=snap.id, mode="as_in_excel",
            scope=scope, include_deleted_lost=include_deleted_lost,
        )
        comp_data = {}
        if comp_snap:
            comp_data = self.analytics_svc.get_comparison(
                ADMIN_CONTEXT, comp_snap.snapshot_date, snap.snapshot_date,
                scope=scope, include_deleted_lost=include_deleted_lost,
            )

        # -------------------------------------------------------------
        # Sheet 1: Expiry Q3 Summary
        # -------------------------------------------------------------
        ws1 = wb.create_sheet(title="Expiry Q3 Summary")
        headers1 = [
            "Service Expiry Period",
            "Forecast Category",
            "Today Amount",
            "Today Count",
            "T-Y Amount",
            "T-Y Count",
            "T-LW Amount",
            "T-LW Count",
        ]
        ws1.append(headers1)
        style_header_row(ws1, len(headers1))

        by_cell = fs.get("by_cell", {})
        by_quarter = fs.get("by_quarter", {})
        totals = fs.get("totals", {})

        q_cats = [
            ("Q1-2026", ["Closed", "Commit", "Best Case", "Pipeline"]),
            ("Q1-2027", ["Closed", "Commit", "Best Case"]),
            ("Q2-2026", ["Closed", "Commit", "Best Case", "Pipeline"]),
            ("Q2-2027", ["Commit"]),
            ("Q3-2026", ["Closed", "Commit", "Best Case", "Pipeline"]),
            ("Q3-2027", ["Closed", "Commit"]),
            ("Q4-2026", ["Closed", "Commit", "Best Case", "Pipeline"]),
            ("Q4-2027", ["Commit"]),
        ]

        r_idx = 2
        for q, cats in q_cats:
            q_d = by_quarter.get(q, {})
            t_q = q_d.get("today", {})
            ty_q = q_d.get("delta_vs_yesterday", {})
            tlw_q = q_d.get("delta_vs_lastweek", {})

            ws1.append([
                q,
                "",
                t_q.get("acv", 0),
                int(t_q.get("count", 0)),
                ty_q.get("acv", 0),
                int(ty_q.get("count", 0)),
                tlw_q.get("acv", 0),
                int(tlw_q.get("count", 0)),
            ])
            for c in range(1, 9):
                cell = ws1.cell(r_idx, c)
                cell.fill = BAND_FILL
                cell.font = BOLD_FONT
                cell.border = THIN_BORDER
            ws1.cell(r_idx, 3).number_format = "$#,##0.00"
            ws1.cell(r_idx, 4).number_format = "#,##0"
            ws1.cell(r_idx, 5).number_format = "$#,##0.00"
            ws1.cell(r_idx, 6).number_format = "#,##0"
            ws1.cell(r_idx, 7).number_format = "$#,##0.00"
            ws1.cell(r_idx, 8).number_format = "#,##0"

            for c, v in [
                (5, ty_q.get("acv", 0)),
                (6, ty_q.get("count", 0)),
                (7, tlw_q.get("acv", 0)),
                (8, tlw_q.get("count", 0)),
            ]:
                if v > 0:
                    ws1.cell(r_idx, c).fill = GREEN_FILL
                    ws1.cell(r_idx, c).font = GREEN_FONT
                elif v < 0:
                    ws1.cell(r_idx, c).fill = RED_FILL
                    ws1.cell(r_idx, c).font = RED_FONT
            r_idx += 1

            for cat in cats:
                c_d = by_cell.get(f"{q}|{cat}", by_cell.get((q, cat), {}))
                t_c = c_d.get("today", {})
                ty_c = c_d.get("delta_vs_yesterday", {})
                tlw_c = c_d.get("delta_vs_lastweek", {})
                ws1.append([
                    "",
                    cat,
                    t_c.get("acv", 0),
                    int(t_c.get("count", 0)),
                    ty_c.get("acv", 0),
                    int(ty_c.get("count", 0)),
                    tlw_c.get("acv", 0),
                    int(tlw_c.get("count", 0)),
                ])
                for c in range(1, 9):
                    ws1.cell(r_idx, c).border = THIN_BORDER
                    ws1.cell(r_idx, c).font = REGULAR_FONT
                ws1.cell(r_idx, 3).number_format = "$#,##0.00"
                ws1.cell(r_idx, 4).number_format = "#,##0"
                ws1.cell(r_idx, 5).number_format = "$#,##0.00"
                ws1.cell(r_idx, 6).number_format = "#,##0"
                ws1.cell(r_idx, 7).number_format = "$#,##0.00"
                ws1.cell(r_idx, 8).number_format = "#,##0"

                for c, v in [
                    (5, ty_c.get("acv", 0)),
                    (6, ty_c.get("count", 0)),
                    (7, tlw_c.get("acv", 0)),
                    (8, tlw_c.get("count", 0)),
                ]:
                    if v > 0:
                        ws1.cell(r_idx, c).fill = GREEN_FILL
                        ws1.cell(r_idx, c).font = GREEN_FONT
                    elif v < 0:
                        ws1.cell(r_idx, c).fill = RED_FILL
                        ws1.cell(r_idx, c).font = RED_FONT
                r_idx += 1

        # Grand Total row
        t_tot = totals.get("today", {})
        ty_tot = totals.get("delta_vs_yesterday", {})
        tlw_tot = totals.get("delta_vs_lastweek", {})
        ws1.append([
            "Grand Total",
            "",
            to_float(t_tot.get("acv")),
            to_int(t_tot.get("count")),
            to_float(ty_tot.get("acv")),
            to_int(ty_tot.get("count")),
            to_float(tlw_tot.get("acv")),
            to_int(tlw_tot.get("count")),
        ])
        for c in range(1, 9):
            cell = ws1.cell(r_idx, c)
            cell.fill = NAVY_FILL
            cell.font = WHITE_BOLD_FONT
            cell.border = THIN_BORDER
        ws1.cell(r_idx, 3).number_format = "$#,##0.00"
        ws1.cell(r_idx, 4).number_format = "#,##0"
        ws1.cell(r_idx, 5).number_format = "$#,##0.00"
        ws1.cell(r_idx, 6).number_format = "#,##0"
        ws1.cell(r_idx, 7).number_format = "$#,##0.00"
        ws1.cell(r_idx, 8).number_format = "#,##0"
        autofit_columns(ws1)

        # -------------------------------------------------------------
        # Sheet 2: Approval Status Summary
        # -------------------------------------------------------------
        ws2 = wb.create_sheet(title="Approval Status Summary")
        headers2 = ["Opportunity Approval Status", "Count", "ACV", "% Count", "% ACV"]
        ws2.append(headers2)
        style_header_row(ws2, len(headers2))

        app_rows = aps.get("rows", [])
        for r in app_rows:
            ws2.append([
                r["approval_status"],
                r["count"],
                r["acv"],
                r["pct_count"] / 100,
                r["pct_acv"] / 100,
            ])
            r_num = ws2.max_row
            ws2.cell(r_num, 2).number_format = "#,##0"
            ws2.cell(r_num, 3).number_format = "$#,##0.00"
            ws2.cell(r_num, 4).number_format = "0.0%"
            ws2.cell(r_num, 5).number_format = "0.0%"
            for c in range(1, 6):
                ws2.cell(r_num, c).border = THIN_BORDER

        # Grand Total row
        ws2.append(["Grand Total", aps.get("total_count", 0), aps.get("total_acv", 0), 1.0, 1.0])
        gt_row = ws2.max_row
        for c in range(1, 6):
            ws2.cell(gt_row, c).fill = NAVY_FILL
            ws2.cell(gt_row, c).font = WHITE_BOLD_FONT
            ws2.cell(gt_row, c).border = THIN_BORDER
        ws2.cell(gt_row, 2).number_format = "#,##0"
        ws2.cell(gt_row, 3).number_format = "$#,##0.00"
        ws2.cell(gt_row, 4).number_format = "0.0%"
        ws2.cell(gt_row, 5).number_format = "0.0%"
        autofit_columns(ws2)

        # -------------------------------------------------------------
        # Sheet 3: Approval Status BU Summary
        # -------------------------------------------------------------
        ws3 = wb.create_sheet(title="Approval Status BU Summary")
        headers3 = [
            "Business Unit",
            "Total Opportunities",
            "Approved",
            "Approved - 2nd",
            "Pending-Approval",
            "Blank",
            "Rejected",
        ]
        ws3.append(headers3)
        style_header_row(ws3, len(headers3))

        for r in bu_aps.get("as_in_excel_rows", []):
            ws3.append([
                r.get("business_unit"),
                r.get("total_count", 0),
                r.get("approved_count", 0) or "-",
                r.get("approved_2nd_count", 0) or "-",
                r.get("pending_count", 0) or "-",
                r.get("blank_count", 0) or "-",
                r.get("rejected_count", 0) or "-",
            ])
            rn = ws3.max_row
            ws3.cell(rn, 2).number_format = "#,##0"
            for c in range(1, 8):
                ws3.cell(rn, c).border = THIN_BORDER

        tot_bu = bu_aps.get("total", {})
        ws3.append([
            "Total",
            tot_bu.get("total_count", 0),
            tot_bu.get("approved_count", 0),
            tot_bu.get("approved_2nd_count", 0),
            tot_bu.get("pending_count", 0),
            tot_bu.get("blank_count", 0),
            tot_bu.get("rejected_count", 0),
        ])
        rn = ws3.max_row
        for c in range(1, 8):
            ws3.cell(rn, c).fill = NAVY_FILL
            ws3.cell(rn, c).font = WHITE_BOLD_FONT
            ws3.cell(rn, c).border = THIN_BORDER
            if c > 1:
                ws3.cell(rn, c).number_format = "#,##0"
        autofit_columns(ws3)

        # -------------------------------------------------------------
        # Sheet 4: Forecast Category Summary
        # -------------------------------------------------------------
        ws4 = wb.create_sheet(title="Forecast Category Summary")
        headers4 = [
            "Forecast Category",
            "Today Count",
            "Today ACV",
            "Yesterday Count",
            "Yesterday ACV",
            "Count Diff",
            "ACV Diff",
        ]
        ws4.append(headers4)
        style_header_row(ws4, len(headers4))

        fc_list = comp_data.get("forecast_category", [])
        for fc in fc_list:
            c_label = "(Blank)" if fc.get("category") == "Blank" else fc.get("category")
            ws4.append([
                c_label,
                fc.get("to_count", 0),
                fc.get("to_acv", 0),
                fc.get("from_count", 0),
                fc.get("from_acv", 0),
                fc.get("count_diff", 0),
                fc.get("acv_diff", 0),
            ])
            rn = ws4.max_row
            ws4.cell(rn, 2).number_format = "#,##0"
            ws4.cell(rn, 3).number_format = "$#,##0.00"
            ws4.cell(rn, 4).number_format = "#,##0"
            ws4.cell(rn, 5).number_format = "$#,##0.00"
            ws4.cell(rn, 6).number_format = "#,##0"
            ws4.cell(rn, 7).number_format = "$#,##0.00"
            for c in range(1, 8):
                ws4.cell(rn, c).border = THIN_BORDER

            if fc.get("count_diff", 0) > 0:
                ws4.cell(rn, 6).fill = GREEN_FILL
                ws4.cell(rn, 6).font = GREEN_FONT
            elif fc.get("count_diff", 0) < 0:
                ws4.cell(rn, 6).fill = RED_FILL
                ws4.cell(rn, 6).font = RED_FONT

            if fc.get("acv_diff", 0) > 0:
                ws4.cell(rn, 7).fill = GREEN_FILL
                ws4.cell(rn, 7).font = GREEN_FONT
            elif fc.get("acv_diff", 0) < 0:
                ws4.cell(rn, 7).fill = RED_FILL
                ws4.cell(rn, 7).font = RED_FONT

        # Total row
        tot_to_cnt = sum(fc.get("to_count", 0) for fc in fc_list)
        tot_to_acv = sum(fc.get("to_acv", 0) for fc in fc_list)
        tot_fr_cnt = sum(fc.get("from_count", 0) for fc in fc_list)
        tot_fr_acv = sum(fc.get("from_acv", 0) for fc in fc_list)
        tot_c_diff = sum(fc.get("count_diff", 0) for fc in fc_list)
        tot_a_diff = sum(fc.get("acv_diff", 0) for fc in fc_list)
        ws4.append(["Total", tot_to_cnt, tot_to_acv, tot_fr_cnt, tot_fr_acv, tot_c_diff, tot_a_diff])
        rn = ws4.max_row
        for c in range(1, 8):
            ws4.cell(rn, c).fill = NAVY_FILL
            ws4.cell(rn, c).font = WHITE_BOLD_FONT
            ws4.cell(rn, c).border = THIN_BORDER
        ws4.cell(rn, 2).number_format = "#,##0"
        ws4.cell(rn, 3).number_format = "$#,##0.00"
        ws4.cell(rn, 4).number_format = "#,##0"
        ws4.cell(rn, 5).number_format = "$#,##0.00"
        ws4.cell(rn, 6).number_format = "#,##0"
        ws4.cell(rn, 7).number_format = "$#,##0.00"
        autofit_columns(ws4)

        # -------------------------------------------------------------
        # Sheet 5: Forecast Movement
        # -------------------------------------------------------------
        ws5 = wb.create_sheet(title="Forecast Movement")
        headers5 = ["Movement", "Count", "ACV Total"]
        ws5.append(headers5)
        style_header_row(ws5, len(headers5))

        for m in comp_data.get("movement", []):
            from_c = m.get("from") or "(Blank)"
            to_c = m.get("to") or "(Blank)"
            ws5.append([f"{from_c} → {to_c}", m.get("count", 0), m.get("acv", 0)])
            rn = ws5.max_row
            ws5.cell(rn, 2).number_format = "#,##0"
            ws5.cell(rn, 3).number_format = "$#,##0.00"
            for c in range(1, 4):
                ws5.cell(rn, c).border = THIN_BORDER
        autofit_columns(ws5)

        # -------------------------------------------------------------
        # Sheet 6: ACV Changes
        # -------------------------------------------------------------
        ws6 = wb.create_sheet(title="ACV Changes")
        headers6 = [
            "Opportunity ID 18 Digit",
            "Opportunity Name",
            "Old ACV",
            "New ACV",
            "ACV Diff",
            "Change Type",
        ]
        ws6.append(headers6)
        style_header_row(ws6, len(headers6))

        for bm in comp_data.get("biggest_movers", []):
            ws6.append([
                bm.get("opportunity_id_18"),
                bm.get("opportunity_name"),
                bm.get("old_acv", 0),
                bm.get("new_acv", 0),
                bm.get("acv_diff", 0),
                bm.get("tag"),
            ])
            rn = ws6.max_row
            ws6.cell(rn, 3).number_format = "$#,##0.00"
            ws6.cell(rn, 4).number_format = "$#,##0.00"
            ws6.cell(rn, 5).number_format = "$#,##0.00"
            for c in range(1, 7):
                ws6.cell(rn, c).border = THIN_BORDER
            if bm.get("acv_diff", 0) > 0:
                ws6.cell(rn, 5).fill = GREEN_FILL
                ws6.cell(rn, 5).font = GREEN_FONT
            elif bm.get("acv_diff", 0) < 0:
                ws6.cell(rn, 5).fill = RED_FILL
                ws6.cell(rn, 5).font = RED_FONT
        autofit_columns(ws6)

        # -------------------------------------------------------------
        # Sheet 7: Top 10 Region Summary
        # -------------------------------------------------------------
        ws7 = wb.create_sheet(title="Top 10 Region Summary")
        headers7 = [
            "Sub-Region",
            "Rank",
            "Opportunity ID 18 Digit",
            "Opportunity Name",
            "Forecast ACV Amount",
        ]
        ws7.append(headers7)
        style_header_row(ws7, len(headers7))

        regions = sorted(list(set(o.sub_region for o in opps if o.sub_region)))
        for reg in regions:
            reg_opps = sorted(
                [o for o in opps if o.sub_region == reg],
                key=lambda o: float(o.forecast_acv_amount or 0),
                reverse=True,
            )[:10]
            for rank, o in enumerate(reg_opps, 1):
                ws7.append([
                    reg,
                    rank,
                    o.opportunity_id_18,
                    o.opportunity_name,
                    float(o.forecast_acv_amount or 0),
                ])
                rn = ws7.max_row
                ws7.cell(rn, 2).number_format = "#,##0"
                ws7.cell(rn, 5).number_format = "$#,##0.00"
                for c in range(1, 6):
                    ws7.cell(rn, c).border = THIN_BORDER
        autofit_columns(ws7)

        # -------------------------------------------------------------
        # Sheet 8: Top 10 Region BU Summary
        # -------------------------------------------------------------
        ws8 = wb.create_sheet(title="Top 10 Region BU Summary")
        headers8 = [
            "Sub-Region",
            "Business Unit",
            "Rank",
            "Opportunity ID 18 Digit",
            "Opportunity Name",
            "Forecast ACV Amount",
        ]
        ws8.append(headers8)
        style_header_row(ws8, len(headers8))

        for reg in regions:
            reg_opps = [o for o in opps if o.sub_region == reg]
            bu_max = {}
            for o in reg_opps:
                bu = o.business_unit_raw or "Unknown"
                val = float(o.forecast_acv_amount or 0)
                if bu not in bu_max or val > bu_max[bu]:
                    bu_max[bu] = val
            leading_bu = max(bu_max.keys(), key=lambda b: bu_max[b]) if bu_max else "None"
            bu_opps = sorted(
                [o for o in reg_opps if (o.business_unit_raw or "Unknown") == leading_bu],
                key=lambda o: float(o.forecast_acv_amount or 0),
                reverse=True,
            )[:10]
            for rank, o in enumerate(bu_opps, 1):
                ws8.append([
                    reg,
                    leading_bu,
                    rank,
                    o.opportunity_id_18,
                    o.opportunity_name,
                    float(o.forecast_acv_amount or 0),
                ])
                rn = ws8.max_row
                ws8.cell(rn, 3).number_format = "#,##0"
                ws8.cell(rn, 6).number_format = "$#,##0.00"
                for c in range(1, 7):
                    ws8.cell(rn, c).border = THIN_BORDER
        autofit_columns(ws8)

        # -------------------------------------------------------------
        # Sheet 9: Change Log
        # -------------------------------------------------------------
        ws9 = wb.create_sheet(title="Change Log")
        headers9 = [
            "Opportunity ID 18 Digit",
            "Opportunity Name",
            "Changed Column",
            "Old Value",
            "New Value",
            "Opportunity Status",
            "Change Type",
        ]
        ws9.append(headers9)
        style_header_row(ws9, len(headers9))

        if comp_snap:
            cls = (
                self.db.query(ChangeLog)
                .filter(
                    ChangeLog.snapshot_from_id == comp_snap.id,
                    ChangeLog.snapshot_to_id == snap.id,
                )
                .all()
            )
            for cl in cls:
                ws9.append([
                    cl.opportunity_id_18,
                    cl.opportunity_name,
                    cl.changed_column,
                    cl.old_value,
                    cl.new_value,
                    cl.opportunity_status,
                    cl.change_type,
                ])
                rn = ws9.max_row
                for c in range(1, 8):
                    ws9.cell(rn, c).border = THIN_BORDER
        autofit_columns(ws9)

        # -------------------------------------------------------------
        # Sheet 10: All Opportunities
        # -------------------------------------------------------------
        ws10 = wb.create_sheet(title="All Opportunities")
        headers10 = [
            "Opportunity ID 18 Digit",
            "Opportunity Name",
            "Account Name",
            "Sub-Region",
            "Country / Territory",
            "Business Unit",
            "Forecast Category",
            "Forecast ACV Amount",
            "Close Date",
            "Probability %",
            "Approval Status",
            "Service Expiry Period",
            "Closing Year",
            "Renewal Category",
            "Opportunity Owner",
        ]
        ws10.append(headers10)
        style_header_row(ws10, len(headers10))

        for o in opps:
            ws10.append([
                o.opportunity_id_18,
                o.opportunity_name,
                o.account_name,
                o.sub_region,
                o.country_territory,
                o.business_unit_raw,
                o.forecast_category,
                float(o.forecast_acv_amount or 0),
                o.close_date.isoformat() if o.close_date else "",
                o.probability_pct,
                o.approval_status,
                o.service_expiry_period,
                o.closing_year,
                o.renewal_category,
                o.opportunity_owner,
            ])
            rn = ws10.max_row
            ws10.cell(rn, 8).number_format = "$#,##0.00"
            if o.probability_pct is not None:
                ws10.cell(rn, 10).number_format = "0.0%"
            for c in range(1, 16):
                ws10.cell(rn, c).border = THIN_BORDER
        autofit_columns(ws10)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        filename = f"Mobileum_Renewals_Daily_Update_{snap.snapshot_date.isoformat()}.xlsx"
        return output, filename
