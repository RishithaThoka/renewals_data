import pandas as pd
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
U='/root/.claude/uploads/736ab44d-f949-5dae-a569-09be824efefc/'
RS=U+'70957310-Renewals_Summary_1.xlsx'; CT=U+'56a513c1-Renewal_Comparison_Tool_11.xlsx'
GREEN=RGBColor(0x92,0xD0,0x6E); RED=RGBColor(0xF4,0xB0,0x84); BAND=RGBColor(0xDD,0xEB,0xF7)
prs=Presentation('/home/claude/Mobileum_Renewals_22slides.pptx')
def setc(cell,text,fill='keep'):
    p=cell.text_frame.paragraphs[0]
    if p.runs:
        p.runs[0].text=text
        for r in p.runs[1:]: r.text=''
    else: p.add_run().text=text
    if fill!='keep':
        if fill is None: cell.fill.background()
        else: cell.fill.solid(); cell.fill.fore_color.rgb=fill
def sfill(v,base=None):
    try: x=float(str(v).replace(',',''))
    except: return base
    return GREEN if x>0 else RED if x<0 else base
def money(v,d=2): return '' if v is None or (isinstance(v,float) and pd.isna(v)) else f"{v:,.{d}f}"
def find(s,n): return [sh for sh in s.shapes if sh.name.startswith(n)]
def title(s): return s.shapes.title.text_frame.text if s.shapes.title is not None else ""
def set_text(shape,text):
    r=shape.text_frame.paragraphs[0].runs; r[0].text=text
    for x in r[1:]: x.text=''

# ---- data
ex=pd.read_excel(RS,sheet_name='Expiry Q3 Summary',header=2) if False else None
raw=pd.read_excel(RS,sheet_name='Expiry Q3 Summary',header=None).dropna(how='all')
raw=raw.iloc[1:]
raw.columns=['q','ct','ta','tc','ya','yc','wa','wc']
ap=pd.read_excel(RS,sheet_name='ApprovalStatus_Summary').iloc[0]
tdata=pd.read_excel(RS,sheet_name='Today_Data')
tdata['acv']=pd.to_numeric(tdata['Forecast ACV Amount'],errors='coerce').fillna(0)
tdata['BU']=tdata['Business Unit'].fillna('')
rA=pd.read_excel(RS,sheet_name='Top 10 Region Summary'); rB=pd.read_excel(RS,sheet_name='Top 10 Region BU Summary')
TF=pd.read_excel(CT,sheet_name='TodayForecastSummary'); MV=pd.read_excel(CT,sheet_name='ForecastMovementSummary')
AC=pd.read_excel(CT,sheet_name='ACVChanges'); AS=pd.read_excel(CT,sheet_name='ApprovalStatusChanges')

for s in prs.slides:
    t=title(s)
    sub=find(s,'Data as of')
    if t.startswith('Renewals Summary'):
        tbl=find(s,'Expiry pivot table')[0].table
        rows=[]; cur=None
        for _,r in raw.iterrows():
            f=lambda v:'' if pd.isna(v) else v
            if r.q=='Grand Total': rows.append(('gt',['Grand Total','',r.ta,r.tc,r.ya,r.yc,r.wa,r.wc]))
            elif isinstance(r.q,str): rows.append(('g',[r.q,'',r.ta,r.tc,r.ya,r.yc,r.wa,r.wc]))
            else: rows.append(('c',['',r.ct,r.ta,r.tc,r.ya,r.yc,r.wa,r.wc]))
        assert len(rows)+1==len(tbl.rows),(len(rows),len(tbl.rows))
        for i,(k,v) in enumerate(rows,1):
            vals=[v[0],v[1],money(v[2]),int(v[3]),money(v[4]),int(v[5]),money(v[6]),int(v[7])]
            for j,x in enumerate(vals):
                base=BAND if k=='g' else None
                fill='keep'
                if k=='gt': fill='keep'
                elif j>=4: fill=sfill(x,base)
                setc(tbl.cell(i,j),str(x),fill)
        g=raw[raw.q=='Grand Total'].iloc[0]
        cards=find(s,'KPI card')
        vals=[f"${g.ta/1e6:,.2f}M  |  {int(g.tc):,}",f"{g.ya/1e6:+,.2f}M".replace('+','+$').replace('-','-$')+f"  |  {int(g.yc):+d}",f"{g.wa/1e6:+,.2f}M".replace('+','+$').replace('-','-$')+f"  |  {int(g.wc):+d}"]
        for c,v in zip(cards,vals): c.text_frame.paragraphs[1].runs[0].text=v
    elif t.startswith('Opportunity Approval'):
        t0=find(s,'Approval status table')[0].table
        order=['Approved','Approved - 2nd','Pending-Approval','Blank','Rejected']
        vals=[int(ap['Total Opportunities'])]+[int(ap[c]) for c in order]
        for j,v in enumerate(vals): setc(t0.cell(1,j),str(v))
        ch=find(s,'Approval chart')[0].chart
        cd=CategoryChartData(); cd.categories=order; cd.add_series('Opportunities',vals[1:]); ch.replace_data(cd)
        o=AS['Old Value'].fillna('').astype(str); n=AS['New Value'].fillna('').astype(str)
        A=['Approved','Approved - 2nd']; P=['Pending-Approval','Pending Approval']
        a=int((n.isin(A)&~o.isin(A)).sum()); p=int((n.isin(P)&~o.isin(P)).sum()); r=int((n=='Rejected').sum())
        kp=find(s,'Key points')[0]; ps=kp.text_frame.paragraphs
        for para,(lab,v) in zip(ps[1:],(('Newly approved',a),('Newly pending approval',p),('Newly rejected',r))): para.runs[0].text=f'{lab}:  {v}'
        print('approval changes',a,p,r)
    elif t.startswith('Forecast Category'):
        tb=find(s,'Forecast category table')[0].table
        TF['Forecast Category']=TF['Forecast Category'].fillna('')
        T=TF.set_index('Forecast Category').loc[['Closed','Commit','Best Case','Pipeline','']].reset_index()
        rows=[[c or '(Blank)',int(x.iloc[1]),money(x['Today ACV'],0),money(x['Yesterday ACV'],0),f"{int(x['Count Diff']):+d}",f"{x['ACV Diff']:+,.0f}"] for _,x in T.iterrows() for c in [x['Forecast Category']]]
        rows.append(['Total',int(T['Today Count'].sum()),money(T['Today ACV'].sum(),0),money(T['Yesterday ACV'].sum(),0),f"{int(T['Count Diff'].sum()):+d}",f"{T['ACV Diff'].sum():+,.0f}"])
        for i,v in enumerate(rows,1):
            last=i==len(rows)
            for j,x in enumerate(v): setc(tb.cell(i,j),str(x),(sfill(x,BAND if last else None) if j>=4 else 'keep'))
        mt=find(s,'Forecast movement table')[0].table
        mv=[(m.replace('->','→').strip()+' (Blank)') if m.strip().endswith('->') else m.replace('->','→').strip() for m in MV['Movement']]
        for i in range(1,len(mt.rows)):
            if i<=len(mv): setc(mt.cell(i,0),mv[i-1]); setc(mt.cell(i,1),str(int(MV['Count'].iloc[i-1])))
            else: setc(mt.cell(i,0),' '); setc(mt.cell(i,1),' ')
    elif t.startswith('Top 10 Forecast ACV'):
        tb=find(s,'ACV change table')[0].table
        A_=AC.copy(); A_['old']=pd.to_numeric(A_['Old Value'],errors='coerce').fillna(0); A_['new']=pd.to_numeric(A_['New Value'],errors='coerce').fillna(0)
        A_['chg']=A_['new']-A_['old']; A_=A_.reindex(A_.chg.abs().sort_values(ascending=False,kind='stable').index).head(10)
        for i,(_,x) in enumerate(A_.iterrows(),1):
            typ='New' if x.old==0 else ('Increase' if x.chg>0 else 'Decrease')
            for j,v in enumerate([i,x['Opportunity Name'],money(x.old,0),money(x.new,0),f"{x.chg:+,.0f}",typ]):
                setc(tb.cell(i,j),str(v),(sfill(v) if j==4 else 'keep'))
    elif t.startswith('Top 10 Opportunities –') :
        reg=t.split('–',1)[1].strip(); g=rA[rA['Sub-Region']==reg]; allr=tdata[tdata['Sub-Region']==reg]
        tb=find(s,'Top 10 table')[0].table
        for i in range(10):
            if i<len(g): x=g.iloc[i]; vals=[i+1,x['Opportunity ID 18 Digit'],x['Opportunity Name'],money(x['Forecast ACV Amount'])]
            else: vals=[i+1,' ',' ',' ']
            for j,v in enumerate(vals): setc(tb.cell(i+1,j),str(v))
        setc(tb.cell(11,3),money(g['Forecast ACV Amount'].sum()))
        set_text(sub[0],f"Data as of 05-Oct-2026   |   {reg} total: ${allr.acv.sum()/1e6:,.2f}M across {len(allr):,} opportunities")
    elif t.startswith('Top 10 by Business Unit'):
        reg=t.split('–',1)[1].strip(); allr=tdata[tdata['Sub-Region']==reg]
        gb=rB[rB['Sub-Region']==reg]; mx=gb.groupby('Business Unit')['Forecast ACV Amount'].max(); bu=mx.idxmax()
        g=gb[gb['Business Unit']==bu].sort_values('Forecast ACV Amount',ascending=False,kind='stable')
        tot=allr[allr.BU==bu]
        tb=find(s,'Top 10 by BU table')[0].table
        for i in range(10):
            if i<len(g): x=g.iloc[i]; vals=[i+1,x['Business Unit'],x['Opportunity ID 18 Digit'],x['Opportunity Name'],money(x['Forecast ACV Amount'])]
            else: vals=[i+1,' ',' ',' ',' ']
            for j,v in enumerate(vals): setc(tb.cell(i+1,j),str(v))
        setc(tb.cell(11,4),money(g['Forecast ACV Amount'].sum()))
        set_text(sub[0],f"Data as of 05-Oct-2026   |   {reg} – {bu}: ${tot.acv.sum()/1e6:,.2f}M across {len(tot):,} opportunities")
prs.save('/home/claude/Renewals_Daily_Update_05-Oct-2026.pptx'); print('saved')
