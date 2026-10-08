import sys, os
from pptx import Presentation

prs = Presentation(r'backend\templates\Mobileum_Renewals_Template.pptx')
s5 = prs.slides[5]
print("Title:", s5.shapes.title.text_frame.text if s5.shapes.title else "")
tbl = [sh for sh in s5.shapes if sh.name == 'Top 10 table'][0].table
for i, r in enumerate(tbl.rows):
    print(f"{i:2d}: {[c.text.strip().replace(chr(10), ' ') for c in r.cells]}")
