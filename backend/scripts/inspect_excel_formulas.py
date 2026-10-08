import openpyxl

wb = openpyxl.load_workbook('data/Renewals Summary 3.xlsx', data_only=False)
ws = wb['Top 10 Region Summary']
print("Row 41 (first row of North America):", [c.value for c in ws[41]])
print("Row 50 (last row of North America):", [c.value for c in ws[50]])
