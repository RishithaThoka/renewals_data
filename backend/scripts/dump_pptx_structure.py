import sys
from pptx import Presentation

prs = Presentation('backend/templates/Mobileum_Renewals_Template.pptx')
lines = [f"Total slides: {len(prs.slides)}"]

for idx, s in enumerate(prs.slides):
    title = s.shapes.title.text if s.shapes.title else "NO TITLE"
    lines.append(f"\n=======================================================")
    lines.append(f"--- Slide {idx+1}: {title} ---")
    lines.append(f"=======================================================")
    for sh in s.shapes:
        info = f"Shape: '{sh.name}' ({sh.shape_type})"
        if sh.has_table:
            t = sh.table
            info += f" [Table {len(t.rows)}x{len(t.columns)}]"
            headers = [t.cell(0, c).text_frame.text.replace('\n', ' ') for c in range(len(t.columns))]
            info += f"\n     Headers: {headers}"
            for r in range(1, min(6, len(t.rows))):
                row_vals = [t.cell(r, c).text_frame.text.replace('\n', ' ') for c in range(len(t.columns))]
                info += f"\n     Row {r}: {row_vals}"
            if len(t.rows) > 6:
                last_vals = [t.cell(len(t.rows)-1, c).text_frame.text.replace('\n', ' ') for c in range(len(t.columns))]
                info += f"\n     Last Row ({len(t.rows)-1}): {last_vals}"
        elif sh.has_chart:
            info += f" [Chart {sh.chart.chart_type}]"
            try:
                info += f" categories: {[c.label for c in sh.chart.plots[0].categories]}"
            except Exception as e:
                info += f" categories err: {e}"
        elif sh.has_text_frame:
            info += f" [Text: {repr(sh.text_frame.text[:120])}]"
        lines.append(f"  {info}")

with open('backend/scripts/pptx_structure.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(lines))
print("Saved to backend/scripts/pptx_structure.txt")
