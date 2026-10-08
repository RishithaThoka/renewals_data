import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation

prs = Presentation('backend/templates/Mobileum_Renewals_Template.pptx')
print(f"Total slides: {len(prs.slides)}")
for idx, s in enumerate(prs.slides):
    title = s.shapes.title.text if s.shapes.title else "NO TITLE"
    print(f"\n--- Slide {idx+1}: {title} ---")
    for sh in s.shapes:
        info = f"Shape: '{sh.name}' ({sh.shape_type})"
        if sh.has_table:
            t = sh.table
            info += f" [Table {len(t.rows)}x{len(t.columns)}]"
            headers = [t.cell(0, c).text_frame.text.replace('\n', ' ') for c in range(len(t.columns))]
            info += f"\n     Headers: {headers}"
            for r in range(1, min(4, len(t.rows))):
                row_vals = [t.cell(r, c).text_frame.text.replace('\n', ' ') for c in range(len(t.columns))]
                info += f"\n     Row {r}: {row_vals}"
        elif sh.has_chart:
            info += f" [Chart {sh.chart.chart_type}]"
            try:
                info += f" categories: {[c.label for c in sh.chart.plots[0].categories]}"
            except Exception as e:
                info += f" categories err: {e}"
        elif sh.has_text_frame:
            info += f" [Text: {sh.text_frame.text[:60].replace(chr(10), ' ')}]"
        print("  ", info)
