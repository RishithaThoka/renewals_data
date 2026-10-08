import sys, os
from datetime import date
from decimal import Decimal
import pandas as pd
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData

sys.path.insert(0, os.path.abspath('.'))
from backend.database import SessionLocal
from backend.services.analytics_service import AnalyticsService, _opps_to_df, EXPIRY_PIVOT_PERIODS, filter_expiry_scope
from backend.services.context import ADMIN_CONTEXT
from backend.models.snapshot import UploadSnapshot
from backend.models.opportunity import Opportunity

GREEN = RGBColor(0x92, 0xD0, 0x6E)
RED = RGBColor(0xF4, 0xB0, 0x84)
BAND = RGBColor(0xDD, 0xEB, 0xF7)

def setc(cell, text, fill='keep'):
    p = cell.text_frame.paragraphs[0]
    if p.runs:
        p.runs[0].text = str(text)
        for r in p.runs[1:]:
            r.text = ''
    else:
        p.add_run().text = str(text)
    if fill != 'keep':
        if fill is None:
            cell.fill.background()
        else:
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill

def sfill(v, base=None):
    try:
        x = float(str(v).replace(',', '').replace('$', '').replace('+', '').strip())
    except Exception:
        return base
    return GREEN if x > 0 else RED if x < 0 else base

def money(v, d=2):
    if v is None:
        return ''
    try:
        x = float(v)
        return f"{x:,.{d}f}"
    except Exception:
        return str(v)

def find(s, n):
    return [sh for sh in s.shapes if sh.name.startswith(n)]

def title(s):
    return s.shapes.title.text_frame.text if s.shapes.title is not None else ""

def set_text(shape, text):
    if shape.text_frame.paragraphs and shape.text_frame.paragraphs[0].runs:
        shape.text_frame.paragraphs[0].runs[0].text = text
        for x in shape.text_frame.paragraphs[0].runs[1:]:
            x.text = ''
    else:
        shape.text_frame.text = text

def test_fill():
    db = SessionLocal()
    svc = AnalyticsService(db)
    
    # Target: Oct 6 snapshot (or active today)
    snap = svc.get_active_snapshot(ADMIN_CONTEXT)
    comp_snap = svc._get_adjacent_snapshot(snap, days_back=1)
    lw_snap = svc._get_adjacent_snapshot(snap, days_back=7)
    
    print(f"Snap: {snap.snapshot_date}, Comp: {comp_snap.snapshot_date}, LW: {lw_snap.snapshot_date}")
    
    prs = Presentation('backend/templates/Mobileum_Renewals_Template.pptx')
    fmt_date = snap.snapshot_date.strftime("%d-%b-%Y")
    
    # Opps & DFs
    opps = svc._opps_for(snap)
    df = _opps_to_df(opps)
    
    comp_opps = svc._opps_for(comp_snap) if comp_snap else []
    comp_df = _opps_to_df(comp_opps) if comp_opps else pd.DataFrame()
    
    # Expiry Summary data
    fs = svc.get_forecast_summary(ADMIN_CONTEXT, snapshot_id=snap.id, compare_to=comp_snap.id if comp_snap else None)
    by_cell = fs.get("by_cell", {})
    by_quarter = fs.get("by_quarter", {})
    totals = fs.get("totals", {})
    
    # Comparison data
    comp_data = svc.get_comparison(ADMIN_CONTEXT, comp_snap.snapshot_date, snap.snapshot_date)
    
    # Approval status data
    aps = svc.get_approval_status(ADMIN_CONTEXT, snapshot_id=snap.id, compare_to=comp_snap.id if comp_snap else None)
    bu_aps = svc.get_approval_by_bu(ADMIN_CONTEXT, snapshot_id=snap.id, mode="as_in_excel")
    
    for s in prs.slides:
        t = title(s)
        sub = find(s, 'Data as of')
        
        # Slide 1: Cover
        cov = find(s, 'Cover date')
        if cov:
            set_text(cov[0], f"Data as of {fmt_date}")
            
        # Slide 2: Expiry
        if t.startswith('Renewals Summary'):
            if sub:
                set_text(sub[0], f"Data as of {fmt_date}   |   T-Y = Today vs Yesterday,  T-LW = Today vs Last Week")
            tbl = find(s, 'Expiry pivot table')[0].table
            
            # Populate exact 32 rows matching template structure
            # Quarters in order: Q1-2026, Q1-2027, Q2-2026, Q2-2027, Q3-2026, Q3-2027, Q4-2026, Q4-2027
            q_cats = [
                ('Q1-2026', ['Closed', 'Commit', 'Best Case', 'Pipeline']),
                ('Q1-2027', ['Closed', 'Commit', 'Best Case']),
                ('Q2-2026', ['Closed', 'Commit', 'Best Case', 'Pipeline']),
                ('Q2-2027', ['Commit']),
                ('Q3-2026', ['Closed', 'Commit', 'Best Case', 'Pipeline']),
                ('Q3-2027', ['Closed', 'Commit']),
                ('Q4-2026', ['Closed', 'Commit', 'Best Case', 'Pipeline']),
                ('Q4-2027', ['Commit']),
            ]
            row_idx = 1
            for q, cats in q_cats:
                # Quarter header row
                q_d = by_quarter.get(q, {})
                t_q = q_d.get('today', {})
                ty_q = q_d.get('delta_vs_yesterday', {})
                tlw_q = q_d.get('delta_vs_lastweek', {})
                q_row_vals = [
                    q, ' ',
                    money(t_q.get('acv', 0)), str(int(t_q.get('count', 0))),
                    money(ty_q.get('acv', 0)), str(int(ty_q.get('count', 0))),
                    money(tlw_q.get('acv', 0)), str(int(tlw_q.get('count', 0))),
                ]
                for j, x in enumerate(q_row_vals):
                    fill = BAND if j < 4 else sfill(x, BAND)
                    setc(tbl.cell(row_idx, j), str(x), fill)
                row_idx += 1
                
                # Category rows
                for c in cats:
                    c_d = by_cell.get(f"{q}|{c}", by_cell.get((q, c), {}))
                    t_c = c_d.get('today', {})
                    ty_c = c_d.get('delta_vs_yesterday', {})
                    tlw_c = c_d.get('delta_vs_lastweek', {})
                    c_row_vals = [
                        ' ', c,
                        money(t_c.get('acv', 0)), str(int(t_c.get('count', 0))),
                        money(ty_c.get('acv', 0)), str(int(ty_c.get('count', 0))),
                        money(tlw_c.get('acv', 0)), str(int(tlw_c.get('count', 0))),
                    ]
                    for j, x in enumerate(c_row_vals):
                        fill = None if j < 4 else sfill(x, None)
                        setc(tbl.cell(row_idx, j), str(x), fill)
                    row_idx += 1
            
            # Grand Total row
            t_tot = totals.get('today', {})
            ty_tot = totals.get('delta_vs_yesterday', {})
            tlw_tot = totals.get('delta_vs_lastweek', {})
            gt_vals = [
                'Grand Total', ' ',
                money(t_tot.get('acv', 0)), str(int(t_tot.get('count', 0))),
                money(ty_tot.get('acv', 0)), str(int(ty_tot.get('count', 0))),
                money(tlw_tot.get('acv', 0)), str(int(tlw_tot.get('count', 0))),
            ]
            for j, x in enumerate(gt_vals):
                setc(tbl.cell(row_idx, j), str(x), 'keep')
            
            # 3 KPI cards
            cards = find(s, 'KPI card')
            t_acv_m = t_tot.get('acv', 0) / 1e6
            ty_acv_m = ty_tot.get('acv', 0) / 1e6
            tlw_acv_m = tlw_tot.get('acv', 0) / 1e6
            
            ty_sign = '+' if ty_acv_m >= 0 else '-'
            tlw_sign = '+' if tlw_acv_m >= 0 else '-'
            
            card_texts = [
                f"Total ACV (Today)\n${t_acv_m:,.2f}M  |  {int(t_tot.get('count', 0)):,}",
                f"vs Yesterday\n{ty_sign}${abs(ty_acv_m):,.2f}M  |  {int(ty_tot.get('count', 0)):+d}",
                f"vs Last Week\n{tlw_sign}${abs(tlw_acv_m):,.2f}M  |  {int(tlw_tot.get('count', 0)):+d}",
            ]
            for c_sh, txt in zip(cards, card_texts):
                set_text(c_sh, txt)

        # Slide 3: Approval Status
        elif t.startswith('Opportunity Approval'):
            if sub:
                set_text(sub[0], f"Data as of {fmt_date}")
            tbl = find(s, 'Approval status table')[0].table
            
            app_map = {r['approval_status']: r for r in aps.get('rows', [])}
            order = ['Approved', 'Approved - 2nd', 'Pending-Approval', 'Blank', 'Rejected']
            counts = [app_map.get(k, {}).get('count', 0) for k in order]
            acvs = [app_map.get(k, {}).get('acv', 0.0) for k in order]
            
            tot_cnt = aps.get('total_count', sum(counts))
            tot_acv = aps.get('total_acv', sum(acvs))
            
            # Row 1: counts
            cnt_row = ['Opportunities', f"{tot_cnt:,}"] + [f"{c:,}" for c in counts]
            for j, val in enumerate(cnt_row):
                setc(tbl.cell(1, j), val)
            
            # Row 2: acv
            acv_row = ['ACV (USD)', f"${tot_acv/1e6:,.2f}M"] + [f"${a/1e6:,.2f}M" for a in acvs]
            for j, val in enumerate(acv_row):
                setc(tbl.cell(2, j), val)
            
            # Chart
            charts = find(s, 'Approval chart')
            if charts:
                ch = charts[0].chart
                cd = CategoryChartData()
                cd.categories = order
                cd.add_series('Opportunities', counts)
                ch.replace_data(cd)
            
            # Today's changes
            chg = aps.get('changes', {})
            kp = find(s, 'Key points')
            if kp:
                ps = kp[0].text_frame.paragraphs
                new_app = chg.get('newly_approved_count', 0)
                new_pend = chg.get('newly_pending_count', 0)
                new_rej = chg.get('newly_rejected_count', 0)
                if len(ps) >= 4:
                    if ps[1].runs: ps[1].runs[0].text = f"Newly approved:  {new_app}"
                    if len(ps[1].runs) > 1:
                        for r in ps[1].runs[1:]: r.text = ''
                    if ps[2].runs: ps[2].runs[0].text = f"Newly pending approval:  {new_pend}"
                    if len(ps[2].runs) > 1:
                        for r in ps[2].runs[1:]: r.text = ''
                    if ps[3].runs: ps[3].runs[0].text = f"Newly rejected:  {new_rej}"
                    if len(ps[3].runs) > 1:
                        for r in ps[3].runs[1:]: r.text = ''

        # Slide 4: Forecast Category
        elif t.startswith('Forecast Category'):
            if sub:
                set_text(sub[0], f"Data as of {fmt_date}")
            tb = find(s, 'Forecast category table')[0].table
            
            fc_list = comp_data.get('forecast_category', [])
            fc_map = {item['category']: item for item in fc_list}
            categories = ['Closed', 'Commit', 'Best Case', 'Pipeline', 'Blank']
            
            tot_today_c = sum(fc_map.get(c, {}).get('to_count', 0) for c in categories)
            tot_today_a = sum(fc_map.get(c, {}).get('to_acv', 0.0) for c in categories)
            tot_yest_a = sum(fc_map.get(c, {}).get('from_acv', 0.0) for c in categories)
            tot_cnt_diff = sum(fc_map.get(c, {}).get('count_diff', 0) for c in categories)
            tot_acv_diff = sum(fc_map.get(c, {}).get('acv_diff', 0.0) for c in categories)
            
            for i, c in enumerate(categories, 1):
                item = fc_map.get(c, {})
                c_label = '(Blank)' if c == 'Blank' else c
                row_vals = [
                    c_label,
                    str(int(item.get('to_count', 0))),
                    money(item.get('to_acv', 0), 0),
                    money(item.get('from_acv', 0), 0),
                    f"{int(item.get('count_diff', 0)):+d}",
                    f"{item.get('acv_diff', 0):+,.0f}",
                ]
                for j, x in enumerate(row_vals):
                    fill = sfill(x, None) if j >= 4 else 'keep'
                    setc(tb.cell(i, j), str(x), fill)
            
            # Total row
            tot_vals = [
                'Total',
                str(int(tot_today_c)),
                money(tot_today_a, 0),
                money(tot_yest_a, 0),
                f"{int(tot_cnt_diff):+d}",
                f"{tot_acv_diff:+,.0f}",
            ]
            for j, x in enumerate(tot_vals):
                fill = sfill(x, BAND) if j >= 4 else BAND
                setc(tb.cell(6, j), str(x), fill)
                
            # Movement table
            mt = find(s, 'Forecast movement table')[0].table
            movements = comp_data.get('movement', [])
            for r_idx in range(1, len(mt.rows)):
                if r_idx - 1 < len(movements):
                    mv = movements[r_idx - 1]
                    from_c = mv.get('from') or '(Blank)'
                    to_c = mv.get('to') or '(Blank)'
                    mv_text = f"{from_c} → {to_c}"
                    setc(mt.cell(r_idx, 0), mv_text)
                    setc(mt.cell(r_idx, 1), str(int(mv.get('count', 0))))
                else:
                    setc(mt.cell(r_idx, 0), ' ')
                    setc(mt.cell(r_idx, 1), ' ')

        # Slide 5: Top 10 ACV Changes
        elif t.startswith('Top 10 Forecast ACV'):
            if sub:
                set_text(sub[0], f"Data as of {fmt_date}")
            tb = find(s, 'ACV change table')[0].table
            movers = comp_data.get('biggest_movers', [])[:10]
            for i in range(10):
                if i < len(movers):
                    m = movers[i]
                    vals = [
                        i + 1,
                        m.get('opportunity_name') or '',
                        money(m.get('old_acv', 0), 0),
                        money(m.get('new_acv', 0), 0),
                        m.get('tag', 'Increase'),
                    ]
                else:
                    vals = [i + 1, ' ', ' ', ' ', ' ']
                for j, v in enumerate(vals):
                    setc(tb.cell(i + 1, j), str(v))

        # Slide 6..14: Top 10 Opportunities – <Region>
        elif t.startswith('Top 10 Opportunities –'):
            reg = t.split('–', 1)[1].strip()
            reg_opps = [o for o in opps if o.sub_region == reg]
            reg_acv = sum(float(o.forecast_acv_amount or 0) for o in reg_opps)
            reg_cnt = len(reg_opps)
            if sub:
                set_text(sub[0], f"Data as of {fmt_date}   |   {reg} total: ${reg_acv/1e6:,.2f}M across {reg_cnt:,} opportunities")
            tb = find(s, 'Top 10 table')[0].table
            top10 = sorted(reg_opps, key=lambda o: float(o.forecast_acv_amount or 0), reverse=True)[:10]
            for i in range(10):
                if i < len(top10):
                    o = top10[i]
                    vals = [i + 1, o.opportunity_id_18 or '', o.opportunity_name or '', money(o.forecast_acv_amount)]
                else:
                    vals = [i + 1, ' ', ' ', ' ']
                for j, v in enumerate(vals):
                    setc(tb.cell(i + 1, j), str(v))
            top10_sum = sum(float(o.forecast_acv_amount or 0) for o in top10)
            setc(tb.cell(11, 3), money(top10_sum))

        # Slide 15: Business Unit – Opportunities & Approval Status
        elif t.startswith('Business Unit –'):
            if sub:
                set_text(sub[0], f"Data as of {fmt_date}   |   Count of opportunities by Business Unit; Blank = no approval status yet")
            tb = find(s, 'BU approval table')[0].table
            bu_rows = bu_aps.get('as_in_excel_rows', [])
            bu_total = bu_aps.get('total', {})
            for i in range(1, len(tb.rows) - 1):
                idx = i - 1
                if idx < len(bu_rows):
                    r = bu_rows[idx]
                    vals = [
                        r.get('business_unit'),
                        f"{r.get('total_count', 0):,}",
                        str(r.get('approved_count', 0)) if r.get('approved_count', 0) > 0 else '–',
                        str(r.get('approved_2nd_count', 0)) if r.get('approved_2nd_count', 0) > 0 else '–',
                        str(r.get('pending_count', 0)) if r.get('pending_count', 0) > 0 else '–',
                        str(r.get('blank_count', 0)) if r.get('blank_count', 0) > 0 else '–',
                        str(r.get('rejected_count', 0)) if r.get('rejected_count', 0) > 0 else '–',
                    ]
                else:
                    vals = [' ', ' ', ' ', ' ', ' ', ' ', ' ']
                for j, v in enumerate(vals):
                    setc(tb.cell(i, j), str(v))
            
            # Total row
            last_r = len(tb.rows) - 1
            tot_vals = [
                'Total',
                f"{bu_total.get('total_count', 0):,}",
                f"{bu_total.get('approved_count', 0):,}",
                f"{bu_total.get('approved_2nd_count', 0):,}",
                f"{bu_total.get('pending_count', 0):,}",
                f"{bu_total.get('blank_count', 0):,}",
                f"{bu_total.get('rejected_count', 0):,}",
            ]
            for j, v in enumerate(tot_vals):
                setc(tb.cell(last_r, j), str(v))

        # Slide 16..24: Top 10 by Business Unit – <Region>
        elif t.startswith('Top 10 by Business Unit –'):
            reg = t.split('–', 1)[1].strip()
            reg_opps = [o for o in opps if o.sub_region == reg]
            
            # leading BU by max single deal
            bu_max = {}
            for o in reg_opps:
                bu = o.business_unit_raw or "Unknown"
                val = float(o.forecast_acv_amount or 0)
                if bu not in bu_max or val > bu_max[bu]:
                    bu_max[bu] = val
            
            leading_bu = max(bu_max.keys(), key=lambda b: bu_max[b]) if bu_max else "None"
            bu_opps = [o for o in reg_opps if (o.business_unit_raw or "Unknown") == leading_bu]
            tot_bu_acv = sum(float(o.forecast_acv_amount or 0) for o in bu_opps)
            tot_bu_cnt = len(bu_opps)
            
            opp_word = 'opportunity' if tot_bu_cnt == 1 else 'opportunities'
            if sub:
                set_text(sub[0], f"Data as of {fmt_date}   |   {reg} – {leading_bu}: ${tot_bu_acv/1e6:,.2f}M across {tot_bu_cnt:,} {opp_word}")
            
            tb = find(s, 'Top 10 by BU table')[0].table
            top10_bu = sorted(bu_opps, key=lambda o: float(o.forecast_acv_amount or 0), reverse=True)[:10]
            for i in range(10):
                if i < len(top10_bu):
                    o = top10_bu[i]
                    vals = [i + 1, leading_bu, o.opportunity_id_18 or '', o.opportunity_name or '', money(o.forecast_acv_amount)]
                else:
                    vals = [i + 1, ' ', ' ', ' ', ' ']
                for j, v in enumerate(vals):
                    setc(tb.cell(i + 1, j), str(v))
            top10_bu_sum = sum(float(o.forecast_acv_amount or 0) for o in top10_bu)
            setc(tb.cell(11, 4), money(top10_bu_sum))

    os.makedirs('data', exist_ok=True)
    out_path = 'data/test_filled.pptx'
    prs.save(out_path)
    print(f"Presentation successfully saved to {out_path}!")
    db.close()

if __name__ == '__main__':
    test_fill()
