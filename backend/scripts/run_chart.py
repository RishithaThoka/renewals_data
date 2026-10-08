from pptx import Presentation
from pptx.chart.data import CategoryChartData

prs = Presentation('backend/templates/Mobileum_Renewals_Template.pptx')
s3 = prs.slides[2]
charts = [sh for sh in s3.shapes if sh.name.startswith('Approval chart')]
if charts:
    ch = charts[0].chart
    cd = CategoryChartData()
    cd.categories = ['Approved', 'Approved - 2nd', 'Pending-Approval', 'Blank', 'Rejected']
    cd.add_series('Opportunities', [748, 877, 31, 1411, 19])
    ch.replace_data(cd)
    print("Chart replace_data succeeded!")
else:
    print("No chart found!")
