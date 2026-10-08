import sys, os
from pptx import Presentation

prs = Presentation('backend/templates/Mobileum_Renewals_Template.pptx')
s14 = prs.slides[14]
print("Title:", s14.shapes.title.text_frame.text if s14.shapes.title else "")
tbl = [sh for sh in s14.shapes if sh.name == 'BU approval table'][0].table
for i, r in enumerate(tbl.rows):
    print(f"{i:2d}: {[c.text.strip().replace(chr(10), ' ') for c in r.cells]}")
