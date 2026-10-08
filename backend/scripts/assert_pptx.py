from pptx import Presentation

prs = Presentation('data/test_filled.pptx')
print(f"Total slides: {len(prs.slides)}")

# Slide 2 assertions
s2 = prs.slides[1]
t2 = [sh for sh in s2.shapes if sh.name.startswith('Expiry pivot table')][0].table
gt_acv = t2.cell(32, 2).text_frame.text
gt_cnt = t2.cell(32, 3).text_frame.text
print(f"Slide 2 Grand Total: ACV={gt_acv}, Count={gt_cnt}")
assert "120,511,647.51" in gt_acv, f"Expected 120,511,647.51 in {gt_acv}"
assert "1167" in gt_cnt or "1,167" in gt_cnt, f"Expected 1167 in {gt_cnt}"

# Slide 3 assertions
s3 = prs.slides[2]
t3 = [sh for sh in s3.shapes if sh.name.startswith('Approval status table')][0].table
app_row = [t3.cell(1, c).text_frame.text for c in range(1, 7)]
print(f"Slide 3 Approval counts: {app_row}")
assert "3,086" in app_row[0]
assert "748" in app_row[1]
assert "877" in app_row[2]
assert "31" in app_row[3]
assert "1,411" in app_row[4]
assert "19" in app_row[5]

# Slide 4 assertions
s4 = prs.slides[3]
t4 = [sh for sh in s4.shapes if sh.name.startswith('Forecast category table')][0].table
fc_rows = {t4.cell(r, 0).text_frame.text: (t4.cell(r, 1).text_frame.text, t4.cell(r, 2).text_frame.text) for r in range(1, 6)}
print(f"Slide 4 Forecast categories: {fc_rows}")
assert fc_rows["Closed"][0] == "1245"
assert fc_rows["Commit"][0] == "488"
assert fc_rows["Best Case"][0] == "547"
assert fc_rows["Pipeline"][0] == "538"
assert fc_rows["(Blank)"][0] == "268"

print("\nALL PPTX ASSERTIONS PASSED PERFECTLY!")
