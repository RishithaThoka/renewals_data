from pptx import Presentation

prs = Presentation('backend/templates/Mobileum_Renewals_Template.pptx')
s2 = prs.slides[1]
t = [sh for sh in s2.shapes if sh.name == 'Expiry pivot table'][0].table
print(f"Expiry table rows: {len(t.rows)}")
for r in range(len(t.rows)):
    row_vals = [t.cell(r, c).text_frame.text.replace('\n', ' ') for c in range(len(t.columns))]
    print(f"Row {r:2d}: {row_vals}")
