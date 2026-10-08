import openpyxl

for fname in ['data/Renewals Summary 3.xlsx', 'data/Renewal Comparison Tool 11.xlsx']:
    print(f"\n==================== {fname} ====================")
    wb = openpyxl.load_workbook(fname, read_only=True)
    print("Sheets:", wb.sheetnames)
    for name in wb.sheetnames:
        ws = wb[name]
        rows = list(ws.iter_rows(max_row=3, values_only=True))
        print(f"\n  --- Sheet: {name} ---")
        for r_idx, r in enumerate(rows):
            print(f"    Row {r_idx+1}: {r[:10]}")
