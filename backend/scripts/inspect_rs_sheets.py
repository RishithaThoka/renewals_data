import openpyxl

wb = openpyxl.load_workbook('data/Renewals Summary 3.xlsx', data_only=True)
for sheet in wb.sheetnames[:7]:
    ws = wb[sheet]
    print(f"\n================ Sheet: {sheet} ================")
    for r in range(1, 10):
        row_vals = [ws.cell(r, c).value for c in range(1, 12)]
        if any(v is not None for v in row_vals):
            print(f"Row {r}: {row_vals}")
